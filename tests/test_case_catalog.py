"""
Tests for pipeline/exporters/case_catalog.py
"""

import json
import sys
from pathlib import Path

import pytest

# Ensure project root is on sys.path
sys.path.insert(0, str(Path(__file__).parent.parent))

from pipeline.exporters.case_catalog import (
    CATALOG_COLUMNS,
    build_catalog,
    normalize_title,
    write_catalog,
    _build_row,
)


# ── normalize_title ────────────────────────────────────────────────────────────

class TestNormalizeTitle:
    def test_simple_lowercases(self):
        assert normalize_title("Ban the Box") == "ban the box"

    def test_strips_trailing_period(self):
        assert normalize_title("Fast Food Co.") == "fast food co"

    def test_handles_hyphens(self):
        # "Co-V(id)accinated" → letters joined via space where punct was
        result = normalize_title("Co-V(id)accinated")
        assert "co" in result
        assert "v" in result
        assert "id" in result
        assert "accinated" in result

    def test_collapses_whitespace(self):
        assert normalize_title("  Pre-K   Education  ") == "pre k education"

    def test_comma_removed(self):
        result = normalize_title("Pay Me My Money, In Cash")
        assert "," not in result
        assert "pay me my money" in result

    def test_empty_string(self):
        assert normalize_title("") == ""

    def test_sparkle_co(self):
        assert normalize_title("Sparkle Co.") == "sparkle co"

    def test_tribeca(self):
        assert normalize_title("TriBeCa Branding") == "tribeca branding"

    def test_race_to_270(self):
        assert normalize_title("Race to 270") == "race to 270"


# ── _build_row ─────────────────────────────────────────────────────────────────

def _make_entry(**kwargs) -> dict:
    base = {
        "case_title": "Test Case",
        "source_pdf": "SomeSchool/Some Book 2020.pdf",
        "source_school": "someschool",
        "source_year": 2020,
        "page_start": 10,
        "page_end": 15,
        "page_count": 6,
        "industry": None,
        "case_type": None,
        "difficulty_overall": None,
        "difficulty_quant": None,
        "difficulty_qual": None,
        "detection_method": "toc",
        "extraction_confidence": 0.85,
        "needs_manual_review": False,
        "output_pdf_path": "cases/SomeSchool/Some Book 2020/test-case.pdf",
    }
    base.update(kwargs)
    return base


class TestBuildRow:
    def test_all_catalog_columns_present(self):
        row = _build_row(_make_entry())
        for col in CATALOG_COLUMNS:
            assert col in row, f"Missing column: {col}"

    def test_normalized_title_generated(self):
        row = _build_row(_make_entry(case_title="Fast Food Co."))
        assert row["normalized_title"] == "fast food co"

    def test_manifest_industry_used_when_no_enrichment(self):
        row = _build_row(_make_entry(industry="Technology"))
        assert row["industry"] == "Technology"

    def test_enrichment_overrides_manifest_for_columbia_2021(self):
        entry = _make_entry(
            case_title="Ban the Box",
            source_pdf="Yale/Columbia 2021.pdf",
            industry=None,
        )
        row = _build_row(entry)
        assert row["industry"] == "Government"
        assert row["case_type"] == "Impact Analysis"
        assert row["difficulty"] == "Medium"

    def test_enrichment_overrides_manifest_for_columbia_2017(self):
        entry = _make_entry(
            case_title="Fast Food Co.",
            source_pdf="Yale/Columbia 2017.pdf",
            industry=None,
        )
        row = _build_row(entry)
        assert row["industry"] == "Food"
        assert row["difficulty_math"] == "Easy"
        assert row["difficulty_structure"] == "Medium"
        assert row["difficulty_creativity"] == "Easy"

    def test_no_enrichment_for_unknown_source(self):
        entry = _make_entry(case_title="Some Random Case", source_pdf="Booth/Booth 2022.pdf")
        row = _build_row(entry)
        # enrichment fields should be None since Booth 2022 has no enrichment
        assert row["difficulty_structure"] is None
        assert row["difficulty_creativity"] is None

    def test_difficulty_maps_from_difficulty_overall(self):
        entry = _make_entry(difficulty_overall="Hard")
        row = _build_row(entry)
        assert row["difficulty"] == "Hard"


# ── build_catalog ──────────────────────────────────────────────────────────────

@pytest.fixture
def tmp_manifest(tmp_path) -> Path:
    """Write a small synthetic manifest and return its path."""
    entries = [
        {
            "id": "aaa",
            "case_title": "Ban the Box",
            "source_pdf": "Yale/Columbia 2021.pdf",
            "source_school": "columbia",
            "source_year": 2021,
            "page_start": 21,
            "page_end": 28,
            "page_count": 8,
            "industry": None,
            "case_type": None,
            "difficulty_overall": None,
            "difficulty_quant": None,
            "difficulty_qual": None,
            "detection_method": "manual_override_columbia_2021",
            "extraction_confidence": 1.0,
            "needs_manual_review": False,
            "output_pdf_path": "cases/Yale/Columbia 2021/ban-the-box.pdf",
        },
        {
            "id": "bbb",
            "case_title": "Cheers",
            "source_pdf": "Yale/Columbia 2021.pdf",
            "source_school": "columbia",
            "source_year": 2021,
            "page_start": 29,
            "page_end": 35,
            "page_count": 7,
            "industry": None,
            "case_type": None,
            "difficulty_overall": None,
            "difficulty_quant": None,
            "difficulty_qual": None,
            "detection_method": "manual_override_columbia_2021",
            "extraction_confidence": 1.0,
            "needs_manual_review": False,
            "output_pdf_path": "cases/Yale/Columbia 2021/cheers.pdf",
        },
        {
            "id": "ccc",
            "case_title": "Some Heuristic Case",
            "source_pdf": "Booth/Booth 2022.pdf",
            "source_school": "booth",
            "source_year": 2022,
            "page_start": 5,
            "page_end": 12,
            "page_count": 8,
            "industry": "Technology",
            "case_type": "Profitability",
            "difficulty_overall": "Medium",
            "difficulty_quant": None,
            "difficulty_qual": None,
            "detection_method": "toc",
            "extraction_confidence": 0.82,
            "needs_manual_review": False,
            "output_pdf_path": "cases/Booth/Booth 2022/some-heuristic-case.pdf",
        },
    ]
    p = tmp_path / "manifest.json"
    p.write_text(json.dumps(entries), encoding="utf-8")
    return p


class TestBuildCatalog:
    def test_returns_correct_count(self, tmp_manifest):
        rows = build_catalog(tmp_manifest)
        assert len(rows) == 3

    def test_columbia_2021_enrichment_applied(self, tmp_manifest):
        rows = build_catalog(tmp_manifest)
        ban = next(r for r in rows if r["case_title"] == "Ban the Box")
        assert ban["industry"] == "Government"
        assert ban["case_type"] == "Impact Analysis"
        assert ban["difficulty"] == "Medium"

    def test_non_enriched_case_uses_manifest_data(self, tmp_manifest):
        rows = build_catalog(tmp_manifest)
        case = next(r for r in rows if r["case_title"] == "Some Heuristic Case")
        assert case["industry"] == "Technology"
        assert case["difficulty"] == "Medium"

    def test_skips_blank_title(self, tmp_path):
        entries = [{"case_title": "", "source_pdf": "X.pdf", "page_start": 1, "page_end": 2}]
        p = tmp_path / "manifest.json"
        p.write_text(json.dumps(entries), encoding="utf-8")
        rows = build_catalog(p)
        assert rows == []

    def test_skips_inverted_page_range(self, tmp_path):
        entries = [{"case_title": "Bad Case", "source_pdf": "X.pdf", "page_start": 10, "page_end": 5}]
        p = tmp_path / "manifest.json"
        p.write_text(json.dumps(entries), encoding="utf-8")
        rows = build_catalog(p)
        assert rows == []

    def test_deduplicates_rows(self, tmp_path):
        entry = {
            "case_title": "Dupe Case", "source_pdf": "X.pdf",
            "page_start": 1, "page_end": 5,
        }
        p = tmp_path / "manifest.json"
        p.write_text(json.dumps([entry, entry]), encoding="utf-8")
        rows = build_catalog(p)
        assert len(rows) == 1

    def test_all_rows_have_catalog_columns(self, tmp_manifest):
        rows = build_catalog(tmp_manifest)
        for row in rows:
            for col in CATALOG_COLUMNS:
                assert col in row


# ── write_catalog ──────────────────────────────────────────────────────────────

class TestWriteCatalog:
    def test_creates_csv_and_xlsx(self, tmp_path, tmp_manifest):
        rows = build_catalog(tmp_manifest)
        csv_p, xlsx_p = write_catalog(rows, tmp_path)
        assert csv_p.exists()
        assert xlsx_p.exists()

    def test_csv_has_correct_columns(self, tmp_path, tmp_manifest):
        import csv
        rows = build_catalog(tmp_manifest)
        csv_p, _ = write_catalog(rows, tmp_path)
        with open(csv_p, encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)
            headers = reader.fieldnames
        assert headers == CATALOG_COLUMNS

    def test_csv_row_count(self, tmp_path, tmp_manifest):
        import csv
        rows = build_catalog(tmp_manifest)
        csv_p, _ = write_catalog(rows, tmp_path)
        with open(csv_p, encoding="utf-8-sig") as f:
            data = list(csv.DictReader(f))
        assert len(data) == 3

    def test_xlsx_sheet_exists(self, tmp_path, tmp_manifest):
        import openpyxl
        rows = build_catalog(tmp_manifest)
        _, xlsx_p = write_catalog(rows, tmp_path)
        wb = openpyxl.load_workbook(str(xlsx_p))
        assert "Cases" in wb.sheetnames

    def test_xlsx_header_is_bold(self, tmp_path, tmp_manifest):
        import openpyxl
        rows = build_catalog(tmp_manifest)
        _, xlsx_p = write_catalog(rows, tmp_path)
        wb = openpyxl.load_workbook(str(xlsx_p))
        ws = wb["Cases"]
        assert ws.cell(row=1, column=1).font.bold is True

    def test_xlsx_row_count(self, tmp_path, tmp_manifest):
        import openpyxl
        rows = build_catalog(tmp_manifest)
        _, xlsx_p = write_catalog(rows, tmp_path)
        wb = openpyxl.load_workbook(str(xlsx_p))
        ws = wb["Cases"]
        # header + data rows
        assert ws.max_row == 4    # 1 header + 3 data

    def test_empty_rows_writes_header_only(self, tmp_path):
        csv_p, xlsx_p = write_catalog([], tmp_path)
        import csv, openpyxl
        with open(csv_p, encoding="utf-8-sig") as f:
            headers = next(csv.reader(f))
        assert headers == CATALOG_COLUMNS
        wb = openpyxl.load_workbook(str(xlsx_p))
        assert wb["Cases"].max_row == 1
