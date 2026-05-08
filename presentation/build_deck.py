"""
Builds `presentation/Case_Repo_Product.pptx` — a 7-slide deck for the
SOM class, focused on the live web product.

Inputs:
  output/case_catalog.csv         — quantitative claims
  presentation/screenshots/*.png  — captured by capture_screenshots.py

Run with:
    python presentation/build_deck.py
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


HERE = Path(__file__).resolve().parent
REPO_ROOT = HERE.parent
SHOTS = HERE / "screenshots"
OUTPUT = HERE / "Case_Repo_Product.pptx"

# Palette — matches the existing _build_deck.py for visual continuity.
NAVY  = RGBColor(0x0E, 0x2A, 0x47)
SLATE = RGBColor(0x33, 0x4E, 0x68)
TEAL  = RGBColor(0x2E, 0x86, 0xAB)
CORAL = RGBColor(0xE0, 0x6C, 0x4C)
SAND  = RGBColor(0xF3, 0xE9, 0xDD)
INK   = RGBColor(0x1A, 0x1A, 0x1A)
GREY  = RGBColor(0x6B, 0x75, 0x80)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


# ── Primitives ────────────────────────────────────────────────────────────────

def _blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _add_rect(slide, left, top, width, height, fill, *, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
    shape.shadow.inherit = False
    return shape


def _add_text(slide, left, top, width, height, text, *,
              size=18, color=INK, bold=False, italic=False,
              align=PP_ALIGN.LEFT, font="Calibri"):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.06)
    tf.margin_right = Inches(0.06)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    lines = text.split("\n") if isinstance(text, str) else list(text)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color
    return tb


def _add_bullets(slide, left, top, width, height, items, *,
                 size=16, color=INK, bullet_color=TEAL, gap_pt=8):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(gap_pt)
        bullet = p.add_run()
        bullet.text = "▸  "
        bullet.font.name = "Calibri"
        bullet.font.size = Pt(size)
        bullet.font.bold = True
        bullet.font.color.rgb = bullet_color
        body = p.add_run()
        body.text = item
        body.font.name = "Calibri"
        body.font.size = Pt(size)
        body.font.color.rgb = color
    return tb


def _header(slide, title, subtitle=None):
    _add_rect(slide, 0, 0, SLIDE_W, Inches(0.08), TEAL)
    _add_text(slide, Inches(0.5), Inches(0.28), Inches(12.5), Inches(0.7),
              title, size=28, bold=True, color=NAVY)
    if subtitle:
        _add_text(slide, Inches(0.5), Inches(0.9), Inches(12.5), Inches(0.4),
                  subtitle, size=15, color=GREY)


def _footer(slide, idx, total):
    _add_text(slide, Inches(0.5), Inches(7.1), Inches(8), Inches(0.3),
              "Case Repo · cases.baglini.co", size=10, color=GREY)
    _add_text(slide, Inches(11.5), Inches(7.1), Inches(1.3), Inches(0.3),
              f"{idx} / {total}", size=10, color=GREY, align=PP_ALIGN.RIGHT)


def _shadowed_image(slide, image_path, left, top, width, *, height=None):
    """Drop an image with a thin border and subtle visual frame."""
    if not image_path.exists():
        # Render a placeholder rectangle so the deck still builds.
        _add_rect(slide, left, top, width, height or Inches(4),
                  SAND, line=GREY)
        _add_text(slide, left + Inches(0.2), top + Inches(0.2),
                  width - Inches(0.4), Inches(0.5),
                  f"[missing screenshot: {image_path.name}]",
                  size=12, color=CORAL, italic=True)
        return None
    if height is not None:
        pic = slide.shapes.add_picture(str(image_path), left, top,
                                       width=width, height=height)
    else:
        pic = slide.shapes.add_picture(str(image_path), left, top, width=width)
    pic.line.color.rgb = GREY
    pic.line.width = Pt(0.5)
    return pic


# ── Slide builders ────────────────────────────────────────────────────────────

def slide_title(prs, _df):
    s = _blank_slide(prs)
    _add_rect(s, 0, 0, SLIDE_W, SLIDE_H, NAVY)
    _add_rect(s, 0, Inches(3.3), SLIDE_W, Inches(0.06), TEAL)
    _add_text(s, Inches(0.8), Inches(2.3), Inches(11.7), Inches(0.6),
              "CASE REPO", size=18, bold=True, color=TEAL)
    _add_text(s, Inches(0.8), Inches(2.65), Inches(11.7), Inches(1.2),
              "A searchable consulting-case library,\nbuilt for Yale SOM.",
              size=44, bold=True, color=WHITE)
    _add_text(s, Inches(0.8), Inches(4.6), Inches(11.7), Inches(0.5),
              "467 cases. 13 schools. One Yale-only login.",
              size=20, color=SAND)
    _add_text(s, Inches(0.8), Inches(6.5), Inches(11.7), Inches(0.4),
              "Dan Baglini  ·  cases.baglini.co  ·  May 2026",
              size=14, color=GREY)


def slide_problem(prs, _df):
    s = _blank_slide(prs)
    _header(s, "The problem",
            "Pre-MBA case prep is fragmented across schools, formats, and PDFs.")

    # Left: the pain
    _add_rect(s, Inches(0.5), Inches(1.5), Inches(6.0), Inches(5.2), SAND)
    _add_text(s, Inches(0.8), Inches(1.65), Inches(5.5), Inches(0.5),
              "What case prep looks like today", size=18, bold=True, color=NAVY)
    _add_bullets(s, Inches(0.8), Inches(2.2), Inches(5.5), Inches(4.5), [
        "13+ casebooks, one PDF per school per year",
        "Each casebook bundles 5–50 cases, every layout slightly different",
        "No way to ask \u201cgive me a Hard market-entry case in healthcare\u201d",
        "No metadata: difficulty, industry, type are inconsistently labelled",
        "Sharing requires emailing 30 MB PDFs to your friends",
    ], size=15, bullet_color=CORAL)

    # Right: the ask
    _add_rect(s, Inches(6.85), Inches(1.5), Inches(6.0), Inches(5.2), WHITE)
    _add_rect(s, Inches(6.85), Inches(1.5), Inches(6.0), Inches(0.08), TEAL)
    _add_text(s, Inches(7.15), Inches(1.65), Inches(5.5), Inches(0.5),
              "What students want", size=18, bold=True, color=NAVY)
    _add_bullets(s, Inches(7.15), Inches(2.2), Inches(5.5), Inches(4.5), [
        "Find any case in seconds, by topic or company type",
        "See difficulty + industry up front, before opening the PDF",
        "One library, every school, one login",
        "Read the PDF in-browser \u2014 no downloads, no Google Drive shuffle",
        "Trustworthy access: Yale-only, verified emails",
    ], size=15)


def slide_browse(prs, _df):
    s = _blank_slide(prs)
    _header(s, "Search and filter, instantly",
            "Free-text title search. Difficulty, industry, type, school filters.")

    _shadowed_image(s, SHOTS / "01_browse.png",
                    Inches(0.5), Inches(1.55), Inches(8.4))

    # Right callouts column
    callout_x = Inches(9.2)
    callout_w = Inches(3.7)
    _add_rect(s, callout_x, Inches(1.55), callout_w, Inches(0.5), TEAL)
    _add_text(s, callout_x, Inches(1.55), callout_w, Inches(0.5),
              "What's on screen", size=14, bold=True, color=WHITE,
              align=PP_ALIGN.CENTER)

    _add_rect(s, callout_x, Inches(2.05), callout_w, Inches(4.6), WHITE,
              line=GREY)
    _add_bullets(s, callout_x + Inches(0.2), Inches(2.2),
                 callout_w - Inches(0.4), Inches(4.4), [
        "Live full-text title search",
        "One-click difficulty chips",
        "467 cases indexed today",
        "Color-coded difficulty: Easy · Medium · Hard",
        "Source school + year + industry on every card",
        "Click any card to open the case",
    ], size=12, gap_pt=6)


def slide_case_detail(prs, _df):
    s = _blank_slide(prs)
    _header(s, "Open a case, read in-browser",
            "Metadata up front. PDF embedded. No downloads required.")

    # Big card crop
    _shadowed_image(s, SHOTS / "06_case_card.png",
                    Inches(0.5), Inches(1.6), Inches(12.3))

    # Below the card: explanation strip
    _add_rect(s, Inches(0.5), Inches(3.4), Inches(12.3), Inches(3.4), WHITE,
              line=GREY)
    _add_text(s, Inches(0.8), Inches(3.55), Inches(11.7), Inches(0.5),
              "Every case page shows the same six things — at a glance.",
              size=16, bold=True, color=NAVY)

    cols = [
        ("Difficulty",
         "Easy / Medium / Hard plus a calibrated\n1\u201310 score."),
        ("Industry",
         "One of 25 canonical buckets,\nLLM-classified + alias-mapped."),
        ("Case type",
         "Profitability, Market Entry, M&A,\nGuesstimate, ..."),
        ("Firm",
         "Tagged where stated; \u201cUndisclosed\u201d\notherwise."),
        ("Pages",
         "Trimmed exactly to the case \u2014 no\nfront-matter, no sponsor pages."),
        ("PDF viewer",
         "Embedded iframe, presigned R2 URL,\n1-hour expiry."),
    ]
    col_w = Inches(1.95)
    gap = Inches(0.05)
    total_w = Inches(0.5) + Inches(0.3)
    x0 = Inches(0.65)
    y0 = Inches(4.2)
    for i, (label, body) in enumerate(cols):
        x = x0 + i * (col_w + gap)
        _add_text(s, x, y0, col_w, Inches(0.4), label,
                  size=12, bold=True, color=TEAL)
        _add_text(s, x, y0 + Inches(0.4), col_w, Inches(2.0), body,
                  size=11, color=SLATE)
    _ = total_w  # keep ruff happy


def slide_by_numbers(prs, df):
    s = _blank_slide(prs)
    _header(s, "By the numbers",
            "Real catalog snapshot, generated from output/case_catalog.csv at build time.")

    # Top KPI row
    n_cases = len(df)
    n_schools = df["source_school"].dropna().nunique()
    n_years = df["source_year"].dropna().nunique()
    diff_cov = (df["difficulty_normalized"].notna() &
                (df["difficulty_normalized"] != "")).sum()
    pct_cov = round(100 * diff_cov / max(n_cases, 1))

    kpis = [
        (str(n_cases), "cases catalogued"),
        (str(n_schools), "source schools"),
        (str(n_years), "years covered"),
        (f"{pct_cov}%", "with calibrated\ndifficulty"),
    ]
    kw = Inches(2.85)
    gap = Inches(0.25)
    total_w = len(kpis) * kw + (len(kpis) - 1) * gap
    left0 = (SLIDE_W - total_w) / 2
    for i, (num, label) in enumerate(kpis):
        x = left0 + i * (kw + gap)
        _add_rect(s, x, Inches(1.5), kw, Inches(1.4), NAVY)
        _add_text(s, x, Inches(1.55), kw, Inches(0.9), num,
                  size=44, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        _add_text(s, x, Inches(2.4), kw, Inches(0.5), label,
                  size=12, color=SAND, align=PP_ALIGN.CENTER)

    # Bottom: two side-by-side charts (industries + schools)
    industry_counts = (df["industry"].fillna("Unspecified")
                       .replace("", "Unspecified")
                       .value_counts().head(8))
    cd1 = CategoryChartData()
    cd1.categories = list(industry_counts.index[::-1])
    cd1.add_series("Cases", [int(v) for v in industry_counts.values[::-1]])
    chart1 = s.shapes.add_chart(
        XL_CHART_TYPE.BAR_CLUSTERED,
        Inches(0.5), Inches(3.2), Inches(6.1), Inches(3.7),
        cd1).chart
    chart1.has_title = True
    chart1.chart_title.text_frame.text = "Top 8 industries"
    for r in chart1.chart_title.text_frame.paragraphs[0].runs:
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = NAVY
    chart1.has_legend = False
    plot1 = chart1.plots[0]
    plot1.has_data_labels = True
    plot1.data_labels.font.size = Pt(10)
    series1 = plot1.series[0]
    series1.format.fill.solid()
    series1.format.fill.fore_color.rgb = TEAL

    school_counts = (df["source_school"].fillna("unknown")
                     .value_counts().head(8))
    cd2 = CategoryChartData()
    cd2.categories = list(school_counts.index[::-1])
    cd2.add_series("Cases", [int(v) for v in school_counts.values[::-1]])
    chart2 = s.shapes.add_chart(
        XL_CHART_TYPE.BAR_CLUSTERED,
        Inches(6.85), Inches(3.2), Inches(6.0), Inches(3.7),
        cd2).chart
    chart2.has_title = True
    chart2.chart_title.text_frame.text = "Top 8 source schools"
    for r in chart2.chart_title.text_frame.paragraphs[0].runs:
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = NAVY
    chart2.has_legend = False
    plot2 = chart2.plots[0]
    plot2.has_data_labels = True
    plot2.data_labels.font.size = Pt(10)
    plot2.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    series2 = plot2.series[0]
    series2.format.fill.solid()
    series2.format.fill.fore_color.rgb = CORAL


def slide_how_it_works(prs, _df):
    s = _blank_slide(prs)
    _header(s, "How it works",
            "Five lifecycle stages, one Postgres row per case.")

    stages = [
        ("Source", "13+ casebook PDFs", NAVY),
        ("Split", "PyMuPDF + per-school parsers", SLATE),
        ("Enrich", "GPT-5 industry / difficulty", CORAL),
        ("Store", "Postgres metadata · R2 PDFs", TEAL),
        ("Serve", "FastAPI + HTMX + Tailwind", NAVY),
    ]
    n = len(stages)
    total_w = Inches(12.3)
    card_w = Inches(2.05)
    gap = (total_w - card_w * n) / (n - 1)
    y = Inches(1.8)
    for i, (title, body, col) in enumerate(stages):
        x = Inches(0.5) + i * (card_w + gap)
        _add_rect(s, x, y, card_w, Inches(0.55), col)
        _add_text(s, x, y, card_w, Inches(0.55), title,
                  size=15, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        _add_rect(s, x, y + Inches(0.55), card_w, Inches(2.0), WHITE,
                  line=GREY)
        _add_text(s, x + Inches(0.15), y + Inches(0.7),
                  card_w - Inches(0.3), Inches(1.8), body,
                  size=12, color=INK, align=PP_ALIGN.CENTER)
        if i < n - 1:
            ax = x + card_w + Inches(0.04)
            ay = y + Inches(1.4)
            arrow = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, ax,
                                       ay - Inches(0.12),
                                       gap - Inches(0.08), Inches(0.25))
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = TEAL
            arrow.line.fill.background()

    # Stack box at the bottom
    _add_rect(s, Inches(0.5), Inches(4.7), Inches(12.3), Inches(2.1), SAND)
    _add_text(s, Inches(0.8), Inches(4.85), Inches(11.7), Inches(0.4),
              "What's running in production",
              size=15, bold=True, color=NAVY)

    chunks = [
        ("Backend",  "Python · FastAPI · psycopg · Argon2"),
        ("Frontend", "Jinja2 · HTMX · Tailwind"),
        ("Data",     "Postgres (Render) · Cloudflare R2"),
        ("Email",    "Resend · DKIM + SPF + DMARC"),
        ("Hosting",  "Render web service · custom HTTPS"),
    ]
    cw = Inches(2.4)
    for i, (label, body) in enumerate(chunks):
        x = Inches(0.7) + i * cw
        _add_text(s, x, Inches(5.4), cw - Inches(0.1), Inches(0.4),
                  label, size=13, bold=True, color=TEAL)
        _add_text(s, x, Inches(5.85), cw - Inches(0.1), Inches(0.9),
                  body, size=12, color=INK)


def slide_closing(prs, _df):
    s = _blank_slide(prs)
    _add_rect(s, 0, 0, SLIDE_W, SLIDE_H, NAVY)
    _add_rect(s, 0, Inches(3.3), SLIDE_W, Inches(0.06), TEAL)
    _add_text(s, Inches(0.8), Inches(2.4), Inches(11.7), Inches(0.7),
              "Live at cases.baglini.co",
              size=42, bold=True, color=WHITE)
    _add_text(s, Inches(0.8), Inches(3.6), Inches(11.7), Inches(0.5),
              "Yale-only sign-up. Verified email required.",
              size=20, color=SAND)

    # What's next, in three short bullets
    _add_text(s, Inches(0.8), Inches(4.7), Inches(11.7), Inches(0.4),
              "What's next", size=15, bold=True, color=TEAL)
    _add_bullets(s, Inches(0.8), Inches(5.1), Inches(11.7), Inches(1.5), [
        "Saved cases / case-prep playlists",
        "Difficulty-matched practice partner finder",
        "Rate limiting + admin tools (already shipping)",
    ], size=15, color=WHITE, bullet_color=CORAL)

    _add_text(s, Inches(0.8), Inches(6.7), Inches(11.7), Inches(0.4),
              "Questions?  ·  dan.baglini@yale.edu",
              size=14, color=TEAL)


# ── Build orchestration ───────────────────────────────────────────────────────

SLIDE_BUILDERS = [
    slide_title,
    slide_problem,
    slide_browse,
    slide_case_detail,
    slide_by_numbers,
    slide_how_it_works,
    slide_closing,
]


def main():
    df = pd.read_csv(REPO_ROOT / "output" / "case_catalog.csv")
    print(f"Loaded {len(df)} cases from output/case_catalog.csv")

    if not SHOTS.exists() or not list(SHOTS.glob("*.png")):
        print(f"\n[warn] No screenshots found in {SHOTS}")
        print("       Run `python presentation/capture_screenshots.py` first.")
        print("       Building deck with placeholders...\n")

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    total = len(SLIDE_BUILDERS)
    for idx, builder in enumerate(SLIDE_BUILDERS, start=1):
        builder(prs, df)
        # Apply footer + page number to every non-title, non-closing slide.
        if idx not in (1, total):
            _footer(prs.slides[-1], idx, total)

    prs.save(str(OUTPUT))
    print(f"\nWrote {OUTPUT}  ({total} slides)")


if __name__ == "__main__":
    main()
